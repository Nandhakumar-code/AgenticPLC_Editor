import pandas as pd
from build_diary import days

# Flatten the 'time' dictionary into separate columns
for day in days:
    day['Time_Research'] = day['time']['Research']
    day['Time_Implementation'] = day['time']['Implementation']
    day['Time_Testing'] = day['time']['Testing']
    day['Time_Documentation'] = day['time']['Documentation']
    del day['time']

df = pd.DataFrame(days)

# Reorder and rename columns for a professional look
columns = {
    'day': 'Day',
    'date': 'Date',
    'hours': 'Total Hours',
    'objective': 'Objective',
    'background': 'Background',
    'research': 'Research Activities',
    'concepts': 'Concepts Learned',
    'analysis': 'Engineering Analysis',
    'implementation': 'Implementation Work',
    'testing': 'Testing',
    'problems': 'Problems Encountered',
    'solutions': 'Solutions Applied',
    'deliverables': 'Deliverables',
    'knowledge': 'Knowledge Gained',
    'next_plan': 'Next Day Plan',
    'Time_Research': 'Research (Hrs)',
    'Time_Implementation': 'Implementation (Hrs)',
    'Time_Testing': 'Testing (Hrs)',
    'Time_Documentation': 'Documentation (Hrs)'
}

df = df.rename(columns=columns)
df = df[list(columns.values())]

# Write to Excel
writer = pd.ExcelWriter('AgenticPLC_Engineering_Diary.xlsx', engine='openpyxl')
df.to_excel(writer, index=False, sheet_name='Project Diary')

# Access the workbook and worksheet to adjust formatting
workbook = writer.book
worksheet = writer.sheets['Project Diary']

from openpyxl.styles import Alignment
for row in worksheet.iter_rows(min_row=1, max_row=worksheet.max_row, min_col=1, max_col=worksheet.max_column):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical='top')

# Set column widths
for column_cells in worksheet.columns:
    worksheet.column_dimensions[column_cells[0].column_letter].width = 45 # default width for text heavy columns
    
# Specifically adjust numeric columns
for col in ['A', 'B', 'C', 'P', 'Q', 'R', 'S']:
    worksheet.column_dimensions[col].width = 15

writer.close()
print("Excel file generated successfully!")
